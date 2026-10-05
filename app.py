import os, requests, streamlit as st, chromadb
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer

st.set_page_config(page_title="StudyMate RAG", page_icon="📚")
st.title("📚 StudyMate RAG")
st.caption("AI-Powered Exam-Oriented Study Assistant")

@st.cache_resource
def load_model():
    return SentenceTransformer("all-MiniLM-L6-v2")

emb = load_model()
db = chromadb.PersistentClient(path="./chroma_db")
col = db.get_or_create_collection("studymate")

if "ready" not in st.session_state:
    st.session_state.ready = False


def read_file(f):
    name = f.name.lower()

    if name.endswith(".pdf"):
        return "\n".join(p.extract_text() or "" for p in PdfReader(f).pages)

    if name.endswith((".txt", ".csv")):
        return f.read().decode("utf-8", errors="ignore")

    if name.endswith(".docx"):
        from docx import Document
        return "\n".join(p.text for p in Document(f).paragraphs if p.text.strip())

    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(f, data_only=True).active
        return "\n".join(
            " ".join(str(c.value) for c in row if c.value is not None)
            for row in ws.iter_rows()
        )

    if name.endswith(".pptx"):
        from pptx import Presentation
        return "\n".join(
            shape.text
            for slide in Presentation(f).slides
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        )

    return ""


def ollama(prompt):
    try:
        r = requests.post(
            "http://localhost:11434/api/generate",
            json={"model": "llama3.2", "prompt": prompt, "stream": False},
            timeout=180
        )
        return r.json().get("response", "❌ No response from Ollama.")
    except Exception as e:
        return f"❌ Ollama Error: {e}"


def search(q, n=5):
    if col.count() == 0:
        return []
    result = col.query(
        query_embeddings=[emb.encode(q).tolist()],
        n_results=n
    )
    return result["documents"][0]


file = st.file_uploader(
    "📂 Upload your study material",
    type=["pdf", "docx", "pptx", "txt", "csv", "xlsx"]
)

if st.button("🚀 Process File") and file:

    text = read_file(file)

    if not text.strip():
        st.error("❌ No readable text found.")
        st.stop()

    chunks = [
        text[i:i + 600]
        for i in range(0, len(text), 500)
        if text[i:i + 600].strip()
    ]

    try:
        db.delete_collection("studymate")
    except:
        pass

    col = db.get_or_create_collection("studymate")
    col.upsert(
        ids=[str(i) for i in range(len(chunks))],
        documents=chunks,
        embeddings=emb.encode(chunks).tolist()
    )

    st.session_state.ready = True
    st.success(f"✅ {file.name} processed successfully!")
    st.info(f"📊 {len(text):,} characters | {len(chunks)} chunks")

    content = text[:25000]

    prompt = f"""
You are StudyMate, an exam preparation assistant.
Analyze this study material and provide:

1. COMPLETE SUMMARY
2. TOPIC-WISE IMPORTANT POINTS
3. TOPIC-WISE IMPORTANT QUESTIONS

For every topic include:
- Important concepts
- Definitions
- Formulas if available
- Examples if available
- Short questions
- 5-mark questions
- 10-mark questions

Use only the given material. Do not invent information.
Use simple language and clear headings.

STUDY MATERIAL:
{content}
"""

    with st.spinner("🧠 Generating study material..."):
        result = ollama(prompt)

    st.subheader("📖 Summary, Important Points & Questions")
    st.markdown(result)


if st.session_state.ready:

    st.divider()
    st.subheader("💬 Ask Questions From Your Material")

    q = st.chat_input("Ask something from your uploaded material...")

    if q:
        context = "\n\n".join(search(q))

        with st.chat_message("user"):
            st.write(q)

        with st.chat_message("assistant"):
            prompt = f"""
Answer the question using ONLY the study material below.

If the answer is not present, say:
"This information is not available in the uploaded material."

Use simple language.

STUDY MATERIAL:
{context}

QUESTION:
{q}
"""
            st.write(ollama(prompt))